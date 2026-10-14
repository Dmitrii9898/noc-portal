"""Выгрузка заявок в Excel и Word."""
from io import BytesIO

from django.http import HttpResponse
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font


def _rows(applications):
    for item in applications:
        target = item.project.title if item.project_id else item.event.title
        yield [item.pk, str(item.user), item.get_kind_display(), target, item.get_status_display(), item.created_at.strftime("%d.%m.%Y")]


def applications_xlsx(applications):
    wb = Workbook()
    ws = wb.active
    ws.title = "Заявки"
    ws.append(["№", "Участник", "Тип", "Объект", "Статус", "Дата"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in _rows(applications):
        ws.append(row)
    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="applications.xlsx"'
    return response


def applications_docx(applications):
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(12)
    doc.add_heading("Реестр заявок НОЦ «Устойчивое развитие»", level=1)
    doc.add_paragraph("Автор портала: Курбатов Дмитрий Евгеньевич.")
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for index, title in enumerate(["№", "Участник", "Тип", "Объект", "Статус"]):
        table.rows[0].cells[index].text = title
    for row in _rows(applications):
        cells = table.add_row().cells
        for index, value in enumerate(row[:5]):
            cells[index].text = str(value)
    buffer = BytesIO()
    doc.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    response["Content-Disposition"] = 'attachment; filename="applications.docx"'
    return response
